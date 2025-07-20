import os
import json
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from dateutil.parser import parse
from collections import deque
import threading


# ---------- ВСТРОЕННЫЕ ФУНКЦИИ КОНВЕРТЕРА -----------------
def load_conversations(json_file):
    with open(json_file, 'r', encoding='utf-8') as f:
        return json.load(f)


def extract_all_conversation_paths(mapping):
    paths = []
    root_node = mapping.get('root', {})
    for child_id in root_node.get('children', []):
        queue = deque()
        queue.append([child_id])
        while queue:
            current_path = queue.popleft()
            last_node = mapping.get(current_path[-1], {})
            if not last_node.get('children'):
                paths.append(current_path)
            else:
                for next_id in last_node['children']:
                    queue.append(current_path + [next_id])
    return paths


def build_messages_from_path(mapping, path):
    messages = []
    for node_id in path:
        node = mapping.get(node_id, {})
        message = node.get('message')
        if message and message.get('content'):
            role = 'assistant' if message.get('model') else 'user'
            messages.append({
                'role': role,
                'content': message['content'],
                'inserted_at': message.get('inserted_at')
            })
    return messages


def save_chat_to_docx(title, messages, filename, out_folder):
    doc = Document()
    doc.add_heading(title, level=1)
    for msg in messages:
        content = msg.get('content', '')
        if not content:
            continue
        paragraph = doc.add_paragraph()
        if msg.get('role') == 'user':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(content)
            run.italic = True
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(content)
        run.font.size = Pt(11)

    safe_filename = "".join(c for c in filename
                            if c.isalnum() or c in (' ', '-', '_', '.')).rstrip()
    doc.save(os.path.join(out_folder, safe_filename))


def format_chat_name(inserted_at, title):
    try:
        dt = parse(inserted_at)
        return f"{dt.year}-{dt.month:02d}-{dt.day:02d}-{title}.docx"
    except Exception:
        dt = datetime.now()
        return f"{dt.year}-{dt.month:02d}-{dt.day:02d}-{title}.docx"


# ---------- ГРАФИЧЕСКИЙ ИНТЕРФЕЙС -----------------
class DeepSeekGUI(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("DeepSeek интерпретатор")
        self.geometry("460x220")
        self.resizable(False, False)

        self.json_path = tk.StringVar()
        self.target_dir = tk.StringVar(value=os.path.abspath("chats"))

        tk.Label(self,
                 text="DeepSeek интерпретатор.\n"
                      "Эта программа интерпретирует файл экспорта чатов DeepSeek "
                      "формата Json в docx файлы веток диалога",
                 wraplength=440, justify="center").pack(pady=8)

        btn_frame = tk.Frame(self)
        btn_frame.pack(pady=4)
        tk.Button(btn_frame, text="Открыть", width=12,
                  command=self.select_json).pack(side="left", padx=8)
        tk.Button(btn_frame, text="Сохранить", width=12,
                  command=self.start_processing).pack(side="left", padx=8)

        # прогресс-бар
        self.progress = ttk.Progressbar(self, mode='indeterminate', length=400)
        self.progress.pack(pady=6)

        self.status = tk.StringVar(value="Выберите JSON-файл")
        tk.Label(self, textvariable=self.status, fg="blue").pack(pady=4)

    # -------------------------------------------------------------
    def select_json(self):
        path = filedialog.askopenfilename(
            title="Выберите файл экспорта DeepSeek (.json)",
            filetypes=[("JSON files", "*.json"), ("All files", "*.*")])
        if path:
            self.json_path.set(path)
            self.status.set(f"Выбран: {os.path.basename(path)}")

    # -------------------------------------------------------------
    def start_processing(self):
        json_file = self.json_path.get()
        if not json_file:
            messagebox.showwarning("Внимание", "Сначала выберите JSON-файл!")
            return

        out_dir = filedialog.askdirectory(title="Выберите папку для сохранения")
        if not out_dir:
            return
        self.target_dir.set(out_dir)

        # запускаем обработку в отдельном потоке, чтобы GUI не завис
        self.progress.start()
        threading.Thread(target=self.run_conversion, daemon=True).start()

    # -------------------------------------------------------------
    def run_conversion(self):
        try:
            json_file = self.json_path.get()
            out_dir = self.target_dir.get()
            chats_dir = os.path.join(out_dir, "chats")
            os.makedirs(chats_dir, exist_ok=True)

            conversations = load_conversations(json_file)
            if not isinstance(conversations, list):
                conversations = [conversations]

            total = 0
            for conv in conversations:
                title = conv.get('title', 'Без названия')
                inserted_at = conv.get('inserted_at', datetime.now().isoformat())
                mapping = conv.get('mapping', {})
                paths = extract_all_conversation_paths(mapping)
                if not paths:
                    continue

                for idx, path in enumerate(paths, 1):
                    messages = build_messages_from_path(mapping, path)
                    if not messages:
                        continue

                    base_name = format_chat_name(inserted_at, title)
                    if len(paths) > 1:
                        base_name = f"{os.path.splitext(base_name)[0]}-ветка{idx}.docx"
                    save_chat_to_docx(title, messages, base_name, chats_dir)
                    total += 1

            self.after(0, self.progress.stop)
            self.after(0, lambda: messagebox.showinfo(
                "Готово", f"Создано {total} файлов в папке\n{chats_dir}"))
            self.after(0, lambda: self.status.set("Готово!"))

        except json.JSONDecodeError:
            self.after(0, self.progress.stop)
            self.after(0, lambda: messagebox.showerror("Ошибка", "Файл не является корректным JSON"))
        except Exception as e:
            self.after(0, self.progress.stop)
            self.after(0, lambda: messagebox.showerror("Ошибка", str(e)))


# ---------- ЗАПУСК ПРИЛОЖЕНИЯ -----------------
if __name__ == "__main__":
    app = DeepSeekGUI()
    app.mainloop()