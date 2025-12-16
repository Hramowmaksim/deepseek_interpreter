import os
import json
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from docx import Document
from docx.shared import Pt, RGBColor  # Импортируем RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from dateutil.parser import parse
import threading
import traceback

class DeepSeekConverter(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("DeepSeek Converter v6.0")
        self.geometry("800x600")
        
        self.json_path = tk.StringVar()
        self.output_dir = tk.StringVar()
        self.is_processing = False
        
        self.setup_ui()
    
    def setup_ui(self):
        # Заголовок
        tk.Label(self, text="DeepSeek Converter", 
                font=("Arial", 18, "bold"), fg="blue").pack(pady=10)
        
        # Выбор файла
        file_frame = tk.LabelFrame(self, text="1. Выберите файл JSON", padx=10, pady=10)
        file_frame.pack(fill="x", padx=20, pady=5)
        
        tk.Label(file_frame, text="Файл:").grid(row=0, column=0, sticky="w", padx=5)
        tk.Entry(file_frame, textvariable=self.json_path, width=60).grid(row=0, column=1, padx=5)
        tk.Button(file_frame, text="Обзор...", command=self.select_file).grid(row=0, column=2, padx=5)
        
        # Выбор папки
        folder_frame = tk.LabelFrame(self, text="2. Выберите папку для сохранения", padx=10, pady=10)
        folder_frame.pack(fill="x", padx=20, pady=5)
        
        tk.Label(folder_frame, text="Папка:").grid(row=0, column=0, sticky="w", padx=5)
        tk.Entry(folder_frame, textvariable=self.output_dir, width=60).grid(row=0, column=1, padx=5)
        tk.Button(folder_frame, text="Обзор...", command=self.select_folder).grid(row=0, column=2, padx=5)
        
        # Кнопки
        button_frame = tk.Frame(self)
        button_frame.pack(pady=15)
        
        self.convert_btn = tk.Button(button_frame, text="Начать конвертацию", 
                                    command=self.start_conversion,
                                    bg="green", fg="white", font=("Arial", 12),
                                    width=20, height=2)
        self.convert_btn.pack()
        
        # Прогресс
        self.progress = ttk.Progressbar(self, mode='indeterminate', length=750)
        self.progress.pack(pady=10)
        
        # Статус
        self.status_var = tk.StringVar(value="Готов к работе")
        tk.Label(self, textvariable=self.status_var, font=("Arial", 10)).pack(pady=5)
        
        # Лог
        log_frame = tk.LabelFrame(self, text="Лог обработки", padx=10, pady=10)
        log_frame.pack(fill="both", expand=True, padx=20, pady=10)
        
        # Текстовое поле с прокруткой
        self.log_text = tk.Text(log_frame, wrap="word", height=15)
        self.log_text.pack(fill="both", expand=True)
        
        scrollbar = tk.Scrollbar(self.log_text)
        scrollbar.pack(side="right", fill="y")
        self.log_text.config(yscrollcommand=scrollbar.set)
        scrollbar.config(command=self.log_text.yview)
        
        # Кнопки управления логом
        log_buttons = tk.Frame(log_frame)
        log_buttons.pack(fill="x", pady=(10, 0))
        
        tk.Button(log_buttons, text="Очистить лог", command=self.clear_log).pack(side="left", padx=5)
    
    def select_file(self):
        filename = filedialog.askopenfilename(
            title="Выберите conversations.json",
            filetypes=[("JSON files", "*.json")]
        )
        if filename:
            self.json_path.set(filename)
            self.log_message(f"Выбран файл: {os.path.basename(filename)}")
    
    def select_folder(self):
        folder = filedialog.askdirectory(title="Выберите папку для сохранения")
        if folder:
            self.output_dir.set(folder)
            self.log_message(f"Папка сохранения: {folder}")
    
    def log_message(self, message):
        """Записывает сообщение в лог"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert("end", f"[{timestamp}] {message}\n")
        self.log_text.see("end")
        self.update_idletasks()
    
    def clear_log(self):
        self.log_text.delete(1.0, "end")
        self.log_message("Лог очищен")
    
    def start_conversion(self):
        if not self.json_path.get():
            messagebox.showerror("Ошибка", "Выберите JSON файл!")
            return
        
        if not self.output_dir.get():
            messagebox.showerror("Ошибка", "Выберите папку для сохранения!")
            return
        
        self.is_processing = True
        self.convert_btn.config(state="disabled", text="Обработка...")
        self.progress.start()
        self.clear_log()
        
        threading.Thread(target=self.process_file, daemon=True).start()
    
    def process_file(self):
        try:
            json_file = self.json_path.get()
            output_dir = self.output_dir.get()
            
            # Создаем папку для экспорта
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            export_dir = os.path.join(output_dir, f"DeepSeek_Export_{timestamp}")
            os.makedirs(export_dir, exist_ok=True)
            
            self.log_message("=" * 50)
            self.log_message("НАЧАЛО КОНВЕРТАЦИИ")
            self.log_message(f"Файл: {os.path.basename(json_file)}")
            self.log_message(f"Папка: {export_dir}")
            self.log_message("=" * 50)
            
            # Загружаем JSON
            self.log_message("Загрузка файла...")
            
            with open(json_file, 'r', encoding='utf-8') as f:
                conversations = json.load(f)
            
            total_conversations = len(conversations)
            self.log_message(f"Загружено диалогов: {total_conversations}")
            
            success_count = 0
            error_count = 0
            
            # Обрабатываем каждый диалог
            for idx, conv in enumerate(conversations):
                conv_num = idx + 1
                
                if conv_num % 10 == 0:
                    self.status_var.set(f"Обработка: {conv_num}/{total_conversations}")
                    self.log_message(f"Обработка: {conv_num}/{total_conversations}")
                
                try:
                    title = conv.get('title', f'Диалог_{conv_num}')
                    self.log_message(f"Диалог {conv_num}: {title}")
                    
                    # Извлекаем сообщения
                    messages = self.extract_messages(conv)
                    
                    if messages:
                        # Сохраняем в DOCX
                        saved = self.save_to_docx(title, messages, idx, export_dir)
                        
                        if saved:
                            success_count += 1
                            self.log_message(f"  ✓ Сохранено {len(messages)} сообщений")
                        else:
                            error_count += 1
                            self.log_message(f"  ✗ Ошибка сохранения")
                    else:
                        self.log_message(f"  ⚠ Нет сообщений")
                        error_count += 1
                
                except Exception as e:
                    error_count += 1
                    self.log_message(f"  ❌ Ошибка: {str(e)}")
            
            # Завершение
            self.log_message("=" * 50)
            self.log_message("ОБРАБОТКА ЗАВЕРШЕНА")
            self.log_message(f"Успешно: {success_count} | Ошибок: {error_count}")
            self.log_message("=" * 50)
            
            result_msg = (f"Конвертация завершена!\n\n"
                         f"Всего диалогов: {total_conversations}\n"
                         f"Успешно сохранено: {success_count}\n"
                         f"С ошибками: {error_count}\n\n"
                         f"Файлы сохранены в:\n{export_dir}")
            
            messagebox.showinfo("Готово", result_msg)
            
        except Exception as e:
            self.log_message(f"ФАТАЛЬНАЯ ОШИБКА: {str(e)}")
            messagebox.showerror("Ошибка", f"Произошла ошибка:\n{str(e)}")
        
        finally:
            self.after(0, self.finish_processing)
    
    def extract_messages(self, conversation):
        """Извлекает сообщения из диалога"""
        messages = []
        mapping = conversation.get('mapping', {})
        
        if not mapping:
            return messages
        
        # Собираем все узлы с сообщениями
        nodes_with_messages = {}
        
        for node_id, node_data in mapping.items():
            if not node_data or not isinstance(node_data, dict):
                continue
            
            message_data = node_data.get('message')
            if not message_data or not isinstance(message_data, dict):
                continue
            
            # Извлекаем контент из fragments
            content = self.extract_content_from_fragments(message_data)
            
            # Если нет fragments, пробуем content напрямую
            if not content:
                content = message_data.get('content', '')
                if isinstance(content, list):
                    # Обрабатываем список content
                    text_parts = []
                    for item in content:
                        if isinstance(item, dict):
                            # Ищем текст в разных ключах
                            for key in ['text', 'content', 'value']:
                                if key in item and isinstance(item[key], str):
                                    text_parts.append(item[key])
                                    break
                        elif isinstance(item, str):
                            text_parts.append(item)
                    content = '\n'.join(text_parts)
            
            if not content:
                continue
            
            # Определяем роль
            role = 'assistant' if message_data.get('model') else 'user'
            
            nodes_with_messages[node_id] = {
                'node_id': node_id,
                'content': content.strip(),
                'role': role,
                'timestamp': message_data.get('inserted_at', ''),
                'parent': node_data.get('parent'),
                'children': node_data.get('children', [])
            }
        
        # Восстанавливаем порядок сообщений
        return self.reconstruct_order(nodes_with_messages)
    
    def extract_content_from_fragments(self, message_data):
        """Извлекает контент из fragments"""
        fragments = message_data.get('fragments', [])
        
        if not isinstance(fragments, list):
            return ""
        
        text_parts = []
        
        for fragment in fragments:
            if isinstance(fragment, dict):
                fragment_content = fragment.get('content', '')
                fragment_type = fragment.get('type', '')
                
                if fragment_content:
                    if fragment_type == 'REQUEST':
                        text_parts.append(f"👤 ВОПРОС:\n{fragment_content}")
                    elif fragment_type == 'RESPONSE':
                        text_parts.append(f"🤖 ОТВЕТ:\n{fragment_content}")
                    else:
                        text_parts.append(fragment_content)
        
        return '\n\n'.join(text_parts) if text_parts else ""
    
    def reconstruct_order(self, nodes_with_messages):
        """Восстанавливает правильный порядок сообщений"""
        messages = []
        
        if not nodes_with_messages:
            return messages
        
        # Простой способ: сортируем по node_id
        sorted_nodes = sorted(nodes_with_messages.items(), 
                             key=lambda x: (int(x[0]) if x[0].isdigit() else 999, x[0]))
        
        for node_id, data in sorted_nodes:
            messages.append({
                'role': data['role'],
                'content': data['content'],
                'timestamp': data['timestamp'],
                'node_id': node_id
            })
        
        return messages
    
    def save_to_docx(self, title, messages, index, output_dir):
        """Сохраняет диалог в DOCX (исправленная версия)"""
        try:
            # Создаем документ
            doc = Document()
            
            # Заголовок
            doc.add_heading(title[:200], level=1)
            
            # Добавляем сообщения
            for i, msg in enumerate(messages):
                content = msg['content'].strip()
                if not content:
                    continue
                
                # Создаем параграф для сообщения
                para = doc.add_paragraph()
                
                # Определяем выравнивание в зависимости от роли
                if msg['role'] == 'user':
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    
                    # Убираем префикс если есть
                    if content.startswith('👤 ВОПРОС:'):
                        content = content[9:].strip()
                    
                    run = para.add_run(content)
                    run.italic = True
                    run.font.size = Pt(11)
                    
                    # Исправляем цвет - используем RGBColor
                    try:
                        run.font.color.rgb = RGBColor(0, 0, 139)  # Темно-синий
                    except:
                        # Если не поддерживается, просто не задаем цвет
                        pass
                    
                else:  # assistant
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    
                    # Убираем префикс если есть
                    if content.startswith('🤖 ОТВЕТ:'):
                        content = content[8:].strip()
                    
                    run = para.add_run(content)
                    run.font.size = Pt(11)
                    
                    # Исправляем цвет - используем RGBColor
                    try:
                        run.font.color.rgb = RGBColor(0, 100, 0)  # Темно-зеленый
                    except:
                        # Если не поддерживается, просто не задаем цвет
                        pass
                
                # Добавляем разделитель между сообщениями
                if i < len(messages) - 1:
                    doc.add_paragraph("-" * 40)
            
            # Создаем имя файла
            safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '-', '_')).strip()
            if not safe_title:
                safe_title = f"dialog_{index+1}"
            
            # Добавляем дату к имени файла
            try:
                if messages and messages[0].get('timestamp'):
                    dt = parse(messages[0]['timestamp'])
                    date_str = dt.strftime("%Y-%m-%d")
                else:
                    date_str = datetime.now().strftime("%Y-%m-%d")
            except:
                date_str = datetime.now().strftime("%Y-%m-%d")
            
            filename = f"{index+1:04d}_{date_str}_{safe_title[:50]}.docx"
            
            # Очистка имени файла от запрещенных символов
            forbidden = ['<', '>', ':', '"', '/', '\\', '|', '?', '*']
            for char in forbidden:
                filename = filename.replace(char, '_')
            
            # Сохраняем файл
            filepath = os.path.join(output_dir, filename)
            doc.save(filepath)
            
            return True
            
        except Exception as e:
            self.log_message(f"  Ошибка сохранения: {str(e)}")
            import traceback
            self.log_message(traceback.format_exc())
            return False
    
    def finish_processing(self):
        """Завершает обработку"""
        self.is_processing = False
        self.convert_btn.config(state="normal", text="Начать конвертацию")
        self.progress.stop()
        self.status_var.set("Готово!")

# Запуск приложения
if __name__ == "__main__":
    app = DeepSeekConverter()
    app.mainloop()